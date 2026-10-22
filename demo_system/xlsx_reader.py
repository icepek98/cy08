from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .input_contract import REQUIRED_COLUMNS, validate_input_path, validate_transaction_row


def _normalize_value(column: str, value: Any) -> str:
    if value is None:
        return ""
    if column == "date" and isinstance(value, (datetime, date)):
        if isinstance(value, datetime):
            value = value.date()
        return value.isoformat()
    if column == "amount" and isinstance(value, (int, float, Decimal)):
        normalized = format(Decimal(str(value)), "f")
        if "." not in normalized:
            return normalized + ".00"
        integer, fractional = normalized.split(".", 1)
        return f"{integer}.{fractional.ljust(2, '0')}"
    return str(value).strip()


def read_xlsx_transactions(file_path: str | Path) -> list[dict[str, str]]:
    path = validate_input_path(file_path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Formato XLSX esperado.")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None or tuple(headers) != REQUIRED_COLUMNS:
            raise ValueError(
                "Cabeçalho inválido: esperado exatamente: "
                + ", ".join(REQUIRED_COLUMNS)
            )

        records = []
        for line_number, values in enumerate(rows, start=2):
            normalized = {
                column: _normalize_value(column, value)
                for column, value in zip(REQUIRED_COLUMNS, values)
            }
            if not any(normalized.values()):
                continue
            try:
                validate_transaction_row(normalized)
            except ValueError as exc:
                raise ValueError(f"Erro na linha {line_number}: {exc}") from exc
            records.append(normalized)
        return records
    finally:
        workbook.close()
